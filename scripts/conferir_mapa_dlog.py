#!/usr/bin/env python3
"""
Conferência do "Mapa comparativo de Preços" gerado pelo painel no modelo da DLOG.

Confere, sem Excel, o que a seção 8.7 da avaliação pede:

  1. Estrutura — abas e ordem, títulos, os 24 cabeçalhos da linha 17 (iguais
     aos da planilha modelo, se informada), mesclagens dos blocos e painéis
     congelados.
  2. Fórmulas — nenhuma aponta para aba inexistente nem para OUTRA aba de
     item; nenhum #N/A, #REF!, #VALUE! ou #DIV/0! nos resultados gravados;
     nenhuma referência a intervalo numa célula só (o "='Item 5 '!B2:F2").
  3. Números — refaz em Python MIN, MAX, AVERAGE, MEDIAN, STDEV (n-1), COUNT e a
     mediana das medianas sobre as linhas com "Compõe = Sim" e compara com o
     resultado gravado de cada fórmula; com --painel, compara também com o JSON
     do ResultadoPesquisa (tolerância de R$ 0,01).
  4. Regras — toda linha "Compõe = Não" tem motivo (Y) e AA vazio; a coluna V
     é data; os preços da aba CMED são números.

Uso:
    python scripts/conferir_mapa_dlog.py --gerado Mapa.xlsx
    python scripts/conferir_mapa_dlog.py --gerado Mapa.xlsx --modelo "Mapa comparativo de Preços - DLOG.xlsx"
    python scripts/conferir_mapa_dlog.py --gerado Mapa.xlsx --painel resultado.json

Sai com código 1 se houver erro (avisos não alteram o código de saída).
"""

from __future__ import annotations

import argparse
import json
import re
import statistics
import sys
from datetime import date, datetime
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    sys.exit("openpyxl ausente: pip install openpyxl")

# Mantidos em sincronia com lib/mapa-dlog/layout.ts
CABECALHO_BASE = [
    "Número do Item da Pesquisa",
    "Fonte da Pesquisa",
    "Número da Compra",
    "Data/Hora da Compra",
    "Data/Hora da Cotação",
    "Anulada/Revogada",
    "Número do Item da Compra",
    "Modalidade",
    "Critério Julgamento Item",
    "Código do Item",
    "Descrição do Item",
    "Unidade de Fornecimento",
    "Quantidade Ofertada",
    "Preço Unitário",
    "Identificação do Fornecedor",
    "Nome do Fornecedor / Identificação da Fonte",
    "Órgão",
    "UASG",
    "Nome UASG",
    "Referência",
    "Compõe a base de dados final",
    "Mês e Ano",
    "Índice de correção acumulado IPCA",
    "Preço unitário corrigido pelo IPCA",
]
TITULOS = {
    "Consolidação": ("A1", "CONSOLIDAÇÃO DE PESQUISA DE PREÇOS DE MEDICAMENTOS"),
    "MEMÓRIA DE CÁLCULO": ("A1", "MEMÓRIA DE CÁLCULO - PESQUISA DE PREÇOS"),
    "Índices IPCA": ("A1", "TABELA DE ÍNDICES IPCA - CORREÇÃO MONETÁRIA"),
    "CMED": ("A1", "Secretaria Executiva - CMED"),
}
MESCLAS_ITEM = ["A1:F1", "B2:F2", "A5:F5", "A11:F11", "B14:F14", "A16:X16"]
PRIMEIRA = 18
TOL = 0.01
ERROS_EXCEL = ("#N/A", "#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!")
RE_ITEM = re.compile(r"^Item \d+$")
RE_REF_ABA = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-zÀ-ÿ_][\wÀ-ÿ.]*))!(\$?[A-Z]{1,3}\$?\d*(?::\$?[A-Z]{1,3}\$?\d*)?)")


class Relatorio:
    def __init__(self) -> None:
        self.erros: list[str] = []
        self.avisos: list[str] = []
        self.ok: list[str] = []

    def erro(self, msg: str) -> None:
        self.erros.append(msg)

    def aviso(self, msg: str) -> None:
        self.avisos.append(msg)

    def certo(self, msg: str) -> None:
        self.ok.append(msg)


def numero(v) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return None


def perto(a: float | None, b: float | None, tol: float = TOL) -> bool:
    if a is None and b is None:
        return True
    if a is None or b is None:
        return False
    return abs(a - b) <= tol


def abas_item(wb) -> list[str]:
    return [n for n in wb.sheetnames if RE_ITEM.match(n)]


# ── 1. Estrutura ──────────────────────────────────────────────────────────────

def conferir_estrutura(wb, modelo, rel: Relatorio) -> None:
    itens = abas_item(wb)
    esperado = ["Consolidação", "MEMÓRIA DE CÁLCULO", *itens, "Índices IPCA", "CMED"]
    nomes = [n for n in wb.sheetnames if n != "Parâmetros"]
    if nomes != esperado:
        rel.erro(f"Ordem das abas: {wb.sheetnames} (esperado {esperado} + Parâmetros)")
    else:
        rel.certo(f"Abas na ordem do modelo ({len(itens)} itens)")
    for n in wb.sheetnames:
        if n != n.strip():
            rel.erro(f"Nome de aba com espaço nas pontas: {n!r}")

    for aba, (cel, texto) in TITULOS.items():
        if aba in wb.sheetnames and wb[aba][cel].value != texto:
            rel.erro(f"{aba}!{cel}: título {wb[aba][cel].value!r}, esperado {texto!r}")

    cabecalho_modelo = CABECALHO_BASE
    if modelo is not None:
        aba_modelo = next((n for n in modelo.sheetnames if n.strip().startswith("Item")), None)
        if aba_modelo:
            cabecalho_modelo = [
                (str(modelo[aba_modelo].cell(17, c).value or "")).strip() for c in range(1, 25)
            ]
            rel.certo(f"Cabeçalhos comparados com a aba {aba_modelo!r} do modelo")
        else:
            rel.aviso("Modelo sem aba 'Item n': cabeçalhos comparados com a lista do painel")

    for n in itens:
        ws = wb[n]
        lidos = [(str(ws.cell(17, c).value or "")).strip() for c in range(1, 25)]
        if lidos != cabecalho_modelo:
            diff = [f"{chr(64 + i + 1)}: {a!r} ≠ {b!r}" for i, (a, b) in enumerate(zip(lidos, cabecalho_modelo)) if a != b]
            rel.erro(f"{n}: cabeçalhos da linha 17 diferem — {'; '.join(diff[:5])}")
        mesclas = {str(m) for m in ws.merged_cells.ranges}
        faltam = [m for m in MESCLAS_ITEM if m not in mesclas]
        if faltam:
            rel.erro(f"{n}: mesclagens ausentes {faltam}")
    if itens:
        rel.certo("Cabeçalhos A–X e mesclagens das abas de item conferidos")

    for aba in ("Consolidação", "MEMÓRIA DE CÁLCULO"):
        if aba in wb.sheetnames:
            pane = wb[aba].freeze_panes
            if pane != "A6":
                rel.erro(f"{aba}: painel congelado em {pane}, esperado A6")


# ── 2. Fórmulas ───────────────────────────────────────────────────────────────

def conferir_formulas(wb_f, wb_v, rel: Relatorio) -> None:
    abas = set(wb_f.sheetnames)
    n_formulas = 0
    for nome in wb_f.sheetnames:
        ws_f, ws_v = wb_f[nome], wb_v[nome]
        for linha in ws_f.iter_rows():
            for cel in linha:
                v = cel.value
                if isinstance(v, str) and v.startswith("="):
                    n_formulas += 1
                    for m in RE_REF_ABA.finditer(v):
                        alvo = (m.group(1) or m.group(2) or "").replace("''", "'")
                        intervalo = m.group(3)
                        if alvo not in abas:
                            rel.erro(f"{nome}!{cel.coordinate}: referência a aba inexistente {alvo!r} em {v}")
                        elif RE_ITEM.match(alvo) and RE_ITEM.match(nome) and alvo != nome:
                            rel.erro(f"{nome}!{cel.coordinate}: aponta para outra aba de item ({alvo}) em {v}")
                        if ":" in intervalo and re.fullmatch(r"=\s*" + re.escape(m.group(0)), v.strip()):
                            rel.erro(f"{nome}!{cel.coordinate}: intervalo numa célula só ({v})")
                    resultado = ws_v[cel.coordinate].value
                    if isinstance(resultado, str) and resultado.strip() in ERROS_EXCEL:
                        rel.erro(f"{nome}!{cel.coordinate}: resultado gravado {resultado} ({v})")
                    # Resultado vazio ("") é gravado como <v></v> e o openpyxl o lê como None:
                    # só avisa quando a fórmula não tem como devolver texto vazio.
                    referencia_simples = re.fullmatch(r"='[^']+'!\$?[A-Z]{1,3}\$?\d+", v.strip())
                    if resultado is None and "HYPERLINK" not in v and '""' not in v and not referencia_simples:
                        rel.aviso(f"{nome}!{cel.coordinate}: fórmula sem resultado gravado ({v[:60]})")
                elif isinstance(v, str) and v.strip() in ERROS_EXCEL:
                    rel.erro(f"{nome}!{cel.coordinate}: valor de erro {v}")
    rel.certo(f"{n_formulas} fórmulas conferidas (referências e resultados gravados)")


# ── 3. Números e 4. Regras ───────────────────────────────────────────────────

def linhas_base(ws_v):
    """Linhas de dados (a partir da 18) até a primeira linha sem fonte em B."""
    r = PRIMEIRA
    while ws_v.cell(r, 2).value not in (None, ""):
        yield r
        r += 1


def estatisticas(valores: list[float]) -> dict:
    if not valores:
        return {"n": 0, "menor": None, "maior": None, "media": None, "mediana": None, "dp": None}
    return {
        "n": len(valores),
        "menor": min(valores),
        "maior": max(valores),
        "media": statistics.fmean(valores),
        "mediana": statistics.median(valores),
        "dp": statistics.stdev(valores) if len(valores) > 1 else None,
    }


def conferir_item(nome: str, ws_v, rel: Relatorio, painel: dict | None) -> None:
    sim: list[float] = []
    por_fonte: dict[str, list[float]] = {}
    for r in linhas_base(ws_v):
        u = ws_v.cell(r, 21).value
        x = numero(ws_v.cell(r, 24).value)
        aa = ws_v.cell(r, 27).value
        ab = ws_v.cell(r, 28).value
        y = ws_v.cell(r, 25).value
        v = ws_v.cell(r, 22).value
        fonte = ws_v.cell(r, 2).value
        if v is not None and not isinstance(v, (datetime, date)):
            rel.erro(f"{nome}!V{r}: 'Mês e Ano' não é data ({v!r})")
        if u == "Sim":
            if x is None:
                rel.erro(f"{nome}!X{r}: preço corrigido ausente numa linha que compõe")
            else:
                sim.append(x)
                if numero(aa) is None or not perto(numero(aa), x, 1e-6):
                    rel.erro(f"{nome}!AA{r}: auxiliar {aa!r} ≠ X {x}")
        elif u == "Não":
            if not y:
                rel.erro(f"{nome}!Y{r}: linha 'Compõe = Não' sem motivo")
            if aa not in (None, ""):
                rel.erro(f"{nome}!AA{r}: linha 'Compõe = Não' entrando na conta ({aa!r})")
        else:
            rel.erro(f"{nome}!U{r}: valor {u!r} (esperado Sim/Não)")
        if ab == "Sim" and x is not None:
            por_fonte.setdefault(str(fonte), []).append(x)

    e = estatisticas(sim)
    gravado = {
        "menor": numero(ws_v["B6"].value),
        "maior": numero(ws_v["B7"].value),
        "media": numero(ws_v["B8"].value),
        "mediana": numero(ws_v["B9"].value),
        "dp": numero(ws_v["E6"].value),
        "n": numero(ws_v["E8"].value),
    }
    for chave in ("menor", "maior", "media", "mediana", "dp"):
        if not perto(e[chave], gravado[chave]):
            rel.erro(f"{nome}: {chave} recalculado {e[chave]} ≠ gravado {gravado[chave]}")
    if (gravado["n"] or 0) != e["n"]:
        rel.erro(f"{nome}: COUNT recalculado {e['n']} ≠ gravado {gravado['n']}")

    # Bloco por fonte (H5…): mediana de cada fonte e mediana das medianas.
    medianas: list[float] = []
    for r in range(2, 16):
        rotulo = ws_v.cell(r, 8).value
        if not rotulo or rotulo in ("Fonte", "Mediana das medianas") or str(rotulo).startswith("CONSOLIDADO"):
            if rotulo == "Mediana das medianas":
                mm = numero(ws_v.cell(r, 12).value)
                calc = statistics.median(medianas) if medianas else None
                if not perto(calc, mm):
                    rel.erro(f"{nome}: mediana das medianas recalculada {calc} ≠ gravada {mm}")
            continue
        valores = por_fonte.get(str(rotulo), [])
        med = statistics.median(valores) if valores else None
        gravada = numero(ws_v.cell(r, 12).value)
        if not perto(med, gravada):
            rel.erro(f"{nome}: mediana da fonte {rotulo} recalculada {med} ≠ gravada {gravada}")
        if med is not None:
            medianas.append(med)

    if painel is not None:
        rec = painel.get("recomendacao", {})
        if not perto(numero(ws_v["B12"].value), rec.get("precoFinal")):
            rel.erro(f"{nome}: B12 {ws_v['B12'].value} ≠ preço final do painel {rec.get('precoFinal')}")
        pr = painel.get("consolidado", {}).get("porRegistro", {})
        for chave, campo in (("media", "media"), ("mediana", "mediana"), ("menor", "menor")):
            if not perto(gravado[chave], pr.get(campo)):
                rel.erro(f"{nome}: {chave} gravado {gravado[chave]} ≠ painel {pr.get(campo)}")
        pf = painel.get("consolidado", {}).get("porFonte", {})
        if medianas and not perto(statistics.median(medianas), pf.get("mediana")):
            rel.erro(f"{nome}: mediana das medianas ≠ painel {pf.get('mediana')}")


def conferir_cmed(ws_v, rel: Relatorio) -> None:
    textos = 0
    for linha in ws_v.iter_rows(min_row=5):
        cab = [c.value for c in linha]
        if len(cab) >= 13 and isinstance(cab[0], str) and cab[3] and str(cab[3]).isdigit():
            for col in (10, 11, 12):  # K, L, M
                v = cab[col]
                if isinstance(v, str) and v.strip():
                    textos += 1
                    rel.erro(f"CMED!{linha[col].coordinate}: preço gravado como texto ({v!r})")
    if not textos:
        rel.certo("Preços da aba CMED gravados como número")


def carregar_painel(caminho: Path | None) -> list[dict | None] | None:
    if caminho is None:
        return None
    dados = json.loads(caminho.read_text(encoding="utf-8"))
    if isinstance(dados, dict):
        dados = [dados]
    return dados


def main() -> int:
    ap = argparse.ArgumentParser(description="Confere o mapa comparativo de preços gerado pelo painel.")
    ap.add_argument("--gerado", required=True, type=Path, help="XLSX gerado pelo painel")
    ap.add_argument("--modelo", type=Path, help="planilha da DLOG usada como modelo (opcional)")
    ap.add_argument("--painel", type=Path, help="JSON do ResultadoPesquisa (um objeto ou uma lista, na ordem dos itens)")
    args = ap.parse_args()

    rel = Relatorio()
    wb_f = load_workbook(args.gerado, data_only=False)
    wb_v = load_workbook(args.gerado, data_only=True)
    modelo = load_workbook(args.modelo, data_only=False) if args.modelo else None
    painel = carregar_painel(args.painel)

    conferir_estrutura(wb_f, modelo, rel)
    conferir_formulas(wb_f, wb_v, rel)
    itens = abas_item(wb_v)
    painel_por_item: list[dict | None] = []
    if painel is not None:
        # Itens não localizados vêm como null no JSON e não têm aba com base de preços.
        painel_por_item = list(painel) + [None] * max(0, len(itens) - len(painel))
    for i, nome in enumerate(itens):
        conferir_item(nome, wb_v[nome], rel, painel_por_item[i] if painel is not None else None)
    rel.certo(f"Estatísticas de {len(itens)} item(ns) recalculadas sobre as linhas 'Compõe = Sim'")
    if "CMED" in wb_v.sheetnames:
        conferir_cmed(wb_v["CMED"], rel)

    for m in rel.ok:
        print(f"  ok     {m}")
    for m in rel.avisos[:50]:
        print(f"  aviso  {m}")
    if len(rel.avisos) > 50:
        print(f"  aviso  … mais {len(rel.avisos) - 50}")
    for m in rel.erros:
        print(f"  ERRO   {m}")
    print(f"\n{len(rel.erros)} erro(s), {len(rel.avisos)} aviso(s).")
    return 1 if rel.erros else 0


if __name__ == "__main__":
    sys.exit(main())
