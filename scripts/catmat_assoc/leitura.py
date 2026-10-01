"""
Leitura dos quatro arquivos de entrada (secao 2 do prompt).

Cabecalhos sao localizados PELO CONTEUDO, nunca pela posicao:
  * GP: a linha que tem "Registro", "CATMAT" e "Qt_Embal";
  * lista mensal CMED: a linha (entre as 80 primeiras) com "REGISTRO" e colunas PF/PMVG;
  * catalogo e extracao de unidades: CSV com separador "@" em cp1252.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from .normaliza import limpar


def _chave(h: object) -> str:
    return re.sub(r"[^A-Z0-9_%]", "", limpar(h).replace(" ", ""))


# ---------------------------------------------------------------------------
# Grande Padrao

GP_CAMPOS = {
    "N": "n", "NO": "n", "Nº": "n", "REGISTRO": "registro", "GEN": "gen", "GEN.": "gen", "ICMS": "icms",
    "CATMAT": "catmat", "DESCRICAO": "descricao", "UNIDADEDEFORNECIMENTO": "unidade", "QT_EMBAL": "qt",
    "XXX": "xxx", "SUBSTANCIA": "substancia", "PRODUTO": "produto", "APRESENTACAO": "apresentacao",
    "CNPJ": "cnpj", "FABRICANTE": "fabricante",
}
STATUS_NUM = "num"


@dataclass
class LinhaGP:
    linha: int  # linha do Excel (1-based)
    n: int | None
    registro: str
    gen: object
    icms: object
    catmat_bruto: object
    descricao: object
    unidade: object
    qt: object
    xxx: object
    substancia: str
    produto: str
    apresentacao: str
    cnpj: str
    fabricante: str

    @property
    def catmat(self) -> int | None:
        c = self.catmat_bruto
        if isinstance(c, bool):
            return None
        if isinstance(c, (int, float)):
            return int(c)
        if isinstance(c, str) and re.fullmatch(r"\s*\d{5,7}\s*", c):
            return int(c)
        return None

    @property
    def status(self) -> str:
        """num · Não tem · Fora · vazio · outro (texto livre, ex.: 'NÃO SERA CRIADO')."""
        if self.catmat is not None:
            return STATUS_NUM
        c = self.catmat_bruto
        if c is None or (isinstance(c, str) and not c.strip()):
            return "vazio"
        lc = limpar(c)
        if lc == "NAO TEM":
            return "Não tem"
        if lc == "FORA":
            return "Fora"
        return "outro"

    @property
    def raiz(self) -> str:
        return self.registro[:9]


@dataclass
class GrandePadrao:
    caminho: Path
    aba: str
    linha_cabecalho: int
    colunas: dict[str, int]  # campo -> indice de coluna (0-based)
    linhas: list[LinhaGP]
    ultima_coluna: int


def ler_gp(caminho: str | Path) -> GrandePadrao:
    import openpyxl

    caminho = Path(caminho)
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=False)
    ws = wb[wb.sheetnames[0]]
    linhas_brutas = list(ws.iter_rows(values_only=True))
    wb.close()
    cab = None
    for i, r in enumerate(linhas_brutas[:50]):
        ks = {_chave(v) for v in r if v is not None}
        if {"REGISTRO", "CATMAT", "QT_EMBAL"} <= ks:
            cab = i
            break
    if cab is None:
        raise SystemExit(f"[ERRO] {caminho.name}: cabeçalho (Registro, CATMAT, Qt_Embal) não encontrado")
    cols: dict[str, int] = {}
    for j, v in enumerate(linhas_brutas[cab]):
        k = _chave(v)
        if k in GP_CAMPOS and GP_CAMPOS[k] not in cols:
            cols[GP_CAMPOS[k]] = j
    for obrig in ("registro", "catmat", "descricao", "unidade", "qt", "substancia", "apresentacao"):
        if obrig not in cols:
            raise SystemExit(f"[ERRO] {caminho.name}: coluna obrigatória ausente: {obrig}")

    def g(r: tuple, campo: str) -> object:
        j = cols.get(campo)
        return r[j] if j is not None and j < len(r) else None

    out: list[LinhaGP] = []
    for i, r in enumerate(linhas_brutas[cab + 1:], start=cab + 2):
        reg = g(r, "registro")
        if reg is None or str(reg).strip() == "":
            continue
        if isinstance(reg, float):
            reg = int(reg)
        n = g(r, "n")
        out.append(LinhaGP(
            linha=i, n=int(n) if isinstance(n, (int, float)) else None, registro=re.sub(r"\D", "", str(reg)),
            gen=g(r, "gen"), icms=g(r, "icms"), catmat_bruto=g(r, "catmat"), descricao=g(r, "descricao"),
            unidade=g(r, "unidade"), qt=g(r, "qt"), xxx=g(r, "xxx"),
            substancia=str(g(r, "substancia") or ""), produto=str(g(r, "produto") or ""),
            apresentacao=str(g(r, "apresentacao") or ""), cnpj=str(g(r, "cnpj") or ""),
            fabricante=str(g(r, "fabricante") or ""),
        ))
    return GrandePadrao(caminho, ws.title, cab + 1, cols, out, max(len(x) for x in linhas_brutas[cab:cab + 5]))


# ---------------------------------------------------------------------------
# Lista mensal CMED

LISTA_CAMPOS = {
    "SUBSTANCIA": "substancia", "PRINCIPIOATIVO": "substancia", "CNPJ": "cnpj", "LABORATORIO": "laboratorio",
    "CODIGOGGREM": "ggrem", "REGISTRO": "registro", "EAN1": "ean1", "PRODUTO": "produto",
    "APRESENTACAO": "apresentacao", "CLASSETERAPEUTICA": "classe",
    "TIPODEPRODUTO(STATUSDOPRODUTO)": "tipo", "TIPODEPRODUTOSTATUSDOPRODUTO": "tipo", "REGIMEDEPRECO": "regime",
    "RESTRICAOHOSPITALAR": "restricao", "CONFAZ87": "confaz87", "ICMS0%": "icms0", "TARJA": "tarja",
}


@dataclass
class LinhaLista:
    registro: str
    substancia: str
    cnpj: str
    laboratorio: str
    ggrem: str
    produto: str
    apresentacao: str
    classe: str
    tipo: str
    regime: str
    restricao: str
    confaz87: str
    icms0: str
    tarja: str


@dataclass
class ListaMensal:
    caminho: Path
    publicada: str | None
    linha_cabecalho: int
    linhas: list[LinhaLista]
    por_registro: dict[str, list[LinhaLista]] = field(default_factory=dict)


def ler_lista(caminho: str | Path) -> ListaMensal:
    import openpyxl

    caminho = Path(caminho)
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    linhas = list(ws.iter_rows(values_only=True))
    wb.close()
    cab = None
    publicada = None
    for i, r in enumerate(linhas[:80]):
        for v in r[:3]:
            if isinstance(v, str) and "PUBLICADA EM" in limpar(v):
                publicada = re.sub(r"\s+", " ", v).strip()
        ks = [_chave(v) for v in r if v is not None]
        if "REGISTRO" in ks and any(k.startswith(("PF", "PMVG")) for k in ks):
            cab = i
            break
    if cab is None:
        raise SystemExit(f"[ERRO] {caminho.name}: cabeçalho (REGISTRO + PF/PMVG) não encontrado nas 80 primeiras linhas")
    cols: dict[str, int] = {}
    for j, v in enumerate(linhas[cab]):
        k = _chave(v)
        if k in LISTA_CAMPOS and LISTA_CAMPOS[k] not in cols:
            cols[LISTA_CAMPOS[k]] = j

    def g(r: tuple, campo: str) -> str:
        j = cols.get(campo)
        v = r[j] if j is not None and j < len(r) else None
        return "" if v is None else str(v).strip()

    out = []
    for r in linhas[cab + 1:]:
        reg = re.sub(r"\D", "", g(r, "registro"))
        if not reg:
            continue
        out.append(LinhaLista(reg, g(r, "substancia"), g(r, "cnpj"), g(r, "laboratorio"), g(r, "ggrem"),
                              g(r, "produto"), g(r, "apresentacao"), g(r, "classe"), g(r, "tipo"), g(r, "regime"),
                              g(r, "restricao"), g(r, "confaz87"), g(r, "icms0"), g(r, "tarja")))
    lm = ListaMensal(caminho, publicada, cab + 1, out)
    for x in out:
        lm.por_registro.setdefault(x.registro, []).append(x)
    return lm


# ---------------------------------------------------------------------------
# Catalogo CATMAT e extracao de unidades

def _ler_csv_arroba(caminho: Path) -> list[list[str]]:
    texto = caminho.read_bytes().decode("cp1252")
    linhas = texto.splitlines()
    return [ln.split("@") for ln in linhas if ln.strip()]


def ler_catalogo(caminho: str | Path) -> dict[int, str]:
    """codigoItem -> descricaoItem (texto oficial, sem alteracao)."""
    caminho = Path(caminho)
    rows = _ler_csv_arroba(caminho)
    cab = [limpar(x) for x in rows[0]]
    if cab[:2] != ["CODIGOITEM", "DESCRICAOITEM"]:
        raise SystemExit(f"[ERRO] {caminho.name}: cabeçalho inesperado {rows[0]}")
    out: dict[int, str] = {}
    for r in rows[1:]:
        cod, desc = r[0], "@".join(r[1:])
        if cod.strip().isdigit():
            out[int(cod)] = desc
    return out


@dataclass
class UnidadeOficial:
    texto: str  # unidadeFornecimentoCapacidade ("FRASCO 100,00 ML", "COMPRIMIDO")
    unidade: str  # unidadeFornecimento ("FRASCO")
    capacidade: float | None
    sigla: str  # siglaUnidadeMedida ("ML")


@dataclass
class SituacaoCatmat:
    situacao: str  # Ativo / Inativo
    descricao: str
    unidades: list[UnidadeOficial]


def ler_unidades(caminho: str | Path) -> dict[int, SituacaoCatmat]:
    caminho = Path(caminho)
    rows = _ler_csv_arroba(caminho)
    cab = [_chave(x) for x in rows[0]]
    ix = {k: i for i, k in enumerate(cab)}
    need = ["CODIGOCATMAT", "DESCRICAOITEM", "SITUACAOCATMAT", "UNIDADEFORNECIMENTOCAPACIDADE",
            "UNIDADEFORNECIMENTO", "CAPACIDADE", "SIGLAUNIDADEMEDIDA"]
    for k in need:
        if k not in ix:
            raise SystemExit(f"[ERRO] {caminho.name}: coluna ausente {k}")
    out: dict[int, SituacaoCatmat] = {}
    for r in rows[1:]:
        if len(r) < len(cab):
            continue
        cod = int(r[ix["CODIGOCATMAT"]])
        s = out.get(cod)
        if s is None:
            s = out[cod] = SituacaoCatmat(r[ix["SITUACAOCATMAT"]].strip(), r[ix["DESCRICAOITEM"]], [])
        txt = r[ix["UNIDADEFORNECIMENTOCAPACIDADE"]].strip()
        if txt and txt != "NA":
            cap = r[ix["CAPACIDADE"]].strip()
            try:
                capf = float(cap.replace(",", ".")) if cap else None
            except ValueError:
                capf = None
            s.unidades.append(UnidadeOficial(txt, r[ix["UNIDADEFORNECIMENTO"]].strip(), capf,
                                             r[ix["SIGLAUNIDADEMEDIDA"]].strip()))
    return out
