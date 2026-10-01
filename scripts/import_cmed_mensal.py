#!/usr/bin/env python3
"""
Importador das listas mensais da CMED (planilhas "NNN.mmm.aa.xlsx") -> Neon.

Cada planilha e a Lista de Precos de Medicamentos de uma competencia
(001 = jan/2017; NNN = (ano - 2017) * 12 + mes). Guarda TODAS as competencias
e TODAS as aliquotas de ICMS publicadas, em tres tabelas (schema public,
modelos em prisma/schema.prisma):

  CmedCompetencia  uma linha por lista carregada (arquivo, sha256, aliquotas,
                   CAP derivado, cabecalho original)
  CmedPrecoVersao  versoes de preco deduplicadas por hash do conteudo: os meses
                   com os mesmos valores apontam para a mesma versao
  CmedVigencia     (competencia, versao) — a matriz de vigencia

Leitura dirigida pelo cabecalho, nunca pela posicao: o layout, os nomes das
colunas e o conjunto de aliquotas mudam entre competencias. As regras de
coluna e de valor estao em scripts/cmed_layout.py (espelho de
lib/cmed-aliquotas.ts). Nenhum preco e derivado por formula: grava-se sempre o
valor publicado.

REPLACE por competencia (idempotente): recarregar um mes substitui so aquele
mes. Um arquivo por transacao; falha num arquivo nao desfaz os anteriores.
Competencia ja carregada com o mesmo sha256 e pulada (use --recarregar para
forcar). So o conteudo das versoes que ainda nao existem no banco e enviado.

As tabelas sao criadas se ausentes, com a MESMA DDL que `prisma db push`
geraria a partir de prisma/schema.prisma — um db:push posterior e no-op.

Uso:
    python scripts/import_cmed_mensal.py --dir cmed_originais_arr --dry-run
    python scripts/import_cmed_mensal.py --dir cmed_originais_arr --confirm
    python scripts/import_cmed_mensal.py --file "037.jan.20.xlsx" --confirm
    python scripts/import_cmed_mensal.py --dir cmed_originais_arr --somente 037-048 --confirm
    python scripts/import_cmed_mensal.py --dir ... --vigencia "vigencia..cmed...xlsx" --dry-run
    python scripts/import_cmed_mensal.py --textos "textos_CMED_original.docx" --confirm
    python scripts/import_cmed_mensal.py --limpar-orfas --confirm
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import statistics
import sys
import time
import zipfile
from collections import Counter, deque
from concurrent.futures import ProcessPoolExecutor
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from xml.etree import ElementTree

sys.path.insert(0, str(Path(__file__).resolve().parent))
from cmed_layout import (  # noqa: E402
    ALIQUOTAS_ICMS,
    BOOLEANOS,
    INDICE_ALIQUOTA,
    LayoutError,
    booleano,
    campo_da_coluna,
    chave,
    coluna_preco,
    preco,
    rotulo_aliquota,
    texto,
)

MESES = ("jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez")
RE_ARQUIVO = re.compile(r"^(\d{3})\.([a-z]{3})\.(\d{2})\.xlsx$", re.I)
LINHAS_CABECALHO = 80
MAX_EXEMPLOS = 8

TEXTOS = (
    "codigoGgrem", "substancia", "cnpj", "laboratorio", "produto", "apresentacao",
    "classeTerapeutica", "tipoProduto", "regimePreco", "tarja", "comercializacao",
)
BOOLS_VERSAO = BOOLEANOS + ("analiseRecursal",)

# Ordem das colunas do COPY (= colunas de "CmedPrecoVersao" menos id/criadoEm).
VERSAO_COLS = (
    "hash", "registro", "codigoGgrem", "ean", "substancia", "cnpj", "laboratorio", "produto",
    "apresentacao", "classeTerapeutica", "tipoProduto", "regimePreco", "tarja", "comercializacao",
    "restricaoHospitalar", "cap", "confaz87", "icms0", "analiseRecursal", "precoLiberado",
    "pf", "pmvg", "extras",
)
TMP_TIPOS = {
    "ean": "TEXT[]", "restricaoHospitalar": "BOOLEAN", "cap": "BOOLEAN", "confaz87": "BOOLEAN",
    "icms0": "BOOLEAN", "analiseRecursal": "BOOLEAN", "precoLiberado": "BOOLEAN",
    "pf": "JSONB", "pmvg": "JSONB", "extras": "JSONB",
}
TABELAS = ("CmedCompetencia", "CmedPrecoVersao", "CmedVigencia")

# DDL identica a de prisma/schema.prisma (models CmedCompetencia, CmedPrecoVersao,
# CmedVigencia) — ver o comentario la.
DDL = """
CREATE TABLE IF NOT EXISTS public."CmedCompetencia" (
    "competencia"      DATE NOT NULL,
    "sequencia"        INTEGER NOT NULL,
    "dataPublicacao"   DATE,
    "arquivo"          TEXT NOT NULL,
    "sha256"           TEXT NOT NULL,
    "linhas"           INTEGER NOT NULL,
    "registros"        INTEGER NOT NULL,
    "aliquotas"        TEXT[],
    "capPercentual"    DOUBLE PRECISION,
    "colunasOriginais" JSONB NOT NULL,
    "importadoEm"      TIMESTAMP(3) NOT NULL DEFAULT CURRENT_TIMESTAMP,
    CONSTRAINT "CmedCompetencia_pkey" PRIMARY KEY ("competencia")
);
CREATE UNIQUE INDEX IF NOT EXISTS "CmedCompetencia_sequencia_key"
    ON public."CmedCompetencia"("sequencia");

CREATE TABLE IF NOT EXISTS public."CmedPrecoVersao" (
    "id"                  SERIAL NOT NULL,
    "hash"                TEXT NOT NULL,
    "registro"            TEXT NOT NULL,
    "codigoGgrem"         TEXT,
    "ean"                 TEXT[],
    "substancia"          TEXT,
    "cnpj"                TEXT,
    "laboratorio"         TEXT,
    "produto"             TEXT,
    "apresentacao"        TEXT,
    "classeTerapeutica"   TEXT,
    "tipoProduto"         TEXT,
    "regimePreco"         TEXT,
    "tarja"               TEXT,
    "comercializacao"     TEXT,
    "restricaoHospitalar" BOOLEAN,
    "cap"                 BOOLEAN,
    "confaz87"            BOOLEAN,
    "icms0"               BOOLEAN,
    "analiseRecursal"     BOOLEAN,
    "precoLiberado"       BOOLEAN NOT NULL DEFAULT false,
    "pf"                  JSONB NOT NULL,
    "pmvg"                JSONB NOT NULL,
    "extras"              JSONB,
    "criadoEm"            TIMESTAMP(3) NOT NULL DEFAULT CURRENT_TIMESTAMP,
    CONSTRAINT "CmedPrecoVersao_pkey" PRIMARY KEY ("id")
);
CREATE UNIQUE INDEX IF NOT EXISTS "CmedPrecoVersao_hash_key"
    ON public."CmedPrecoVersao"("hash");
CREATE INDEX IF NOT EXISTS "CmedPrecoVersao_registro_idx"
    ON public."CmedPrecoVersao"("registro");

CREATE TABLE IF NOT EXISTS public."CmedVigencia" (
    "competencia" DATE NOT NULL,
    "versaoId"    INTEGER NOT NULL,
    CONSTRAINT "CmedVigencia_pkey" PRIMARY KEY ("competencia","versaoId")
);
CREATE INDEX IF NOT EXISTS "CmedVigencia_versaoId_idx"
    ON public."CmedVigencia"("versaoId");

DO $$
BEGIN
    IF NOT EXISTS (
        SELECT 1 FROM pg_constraint WHERE conname = 'CmedVigencia_versaoId_fkey'
    ) THEN
        ALTER TABLE public."CmedVigencia"
            ADD CONSTRAINT "CmedVigencia_versaoId_fkey"
            FOREIGN KEY ("versaoId") REFERENCES public."CmedPrecoVersao"("id")
            ON DELETE CASCADE ON UPDATE CASCADE;
    END IF;
END
$$;
"""


# ─── leitura de uma lista ────────────────────────────────────────────────────


@dataclass
class Lista:
    """Uma planilha lida e depurada, pronta para a carga."""

    arquivo: str
    nome: str
    sequencia: int = 0
    competencia: date | None = None
    sha256: str = ""
    aba: str = ""
    linha_cabecalho: int = 0
    colunas: list[dict] = field(default_factory=list)
    linhas: int = 0
    declaradas: int | None = None
    registros: int = 0
    aliquotas: list[str] = field(default_factory=list)
    colunas_vazias: list[str] = field(default_factory=list)
    cap_percentual: float | None = None
    versoes: list[tuple] = field(default_factory=list)
    tamanhos: list[int] = field(default_factory=list)
    registros_set: set[str] = field(default_factory=set)
    stats: Counter = field(default_factory=Counter)
    extras: Counter = field(default_factory=Counter)
    exemplos: dict[str, list[str]] = field(default_factory=dict)
    avisos: list[str] = field(default_factory=list)
    segundos: float = 0.0
    erro: str | None = None

    def exemplo(self, tipo: str, valor: str) -> None:
        lst = self.exemplos.setdefault(tipo, [])
        if len(lst) < MAX_EXEMPLOS and valor not in lst:
            lst.append(valor)


def competencia_do_nome(nome: str) -> tuple[int, date, str]:
    """'037.jan.20.xlsx' -> (37, date(2020, 1, 1), 'jan.20'); valida o NNN."""
    m = RE_ARQUIVO.match(nome)
    if not m:
        raise LayoutError(f"nome fora do padrao NNN.mmm.aa.xlsx: {nome}")
    seq, mes_txt, aa = int(m.group(1)), m.group(2).lower(), int(m.group(3))
    if mes_txt not in MESES:
        raise LayoutError(f"mes desconhecido no nome do arquivo: {mes_txt!r}")
    mes, ano = MESES.index(mes_txt) + 1, 2000 + aa
    esperado = (ano - 2017) * 12 + mes
    if seq != esperado:
        raise LayoutError(f"sequencia {seq:03d} nao bate com {mes_txt}/{ano} (esperado {esperado:03d})")
    return seq, date(ano, mes, 1), f"{mes_txt}.{aa:02d}"


def sha256_arquivo(caminho: str) -> str:
    h = hashlib.sha256()
    with open(caminho, "rb") as f:
        for bloco in iter(lambda: f.read(1 << 20), b""):
            h.update(bloco)
    return h.hexdigest()


def _letra(j: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(j + 1)


def _mesclar(a: dict, b: dict) -> dict:
    """Duplicata com os mesmos precos: OR nos booleanos, 1o nao vazio nos textos."""
    r = dict(a)
    for c in BOOLS_VERSAO:
        va, vb = a[c], b[c]
        r[c] = True if (va or vb) else (False if (va is False or vb is False) else None)
    for c in TEXTOS:
        if r[c] is None:
            r[c] = b[c]
    r["ean"] = a["ean"] + [e for e in b["ean"] if e not in a["ean"]]
    r["extras"] = {**b["extras"], **a["extras"]}
    return r


def _canonico(rec: dict) -> str:
    """Conteudo que define a versao. Precos por codigo (nao por posicao) e com
    2 casas: acrescentar aliquota no fim de ALIQUOTAS_ICMS nao muda o hash."""
    d = {k: rec[k] for k in VERSAO_COLS if k not in ("hash", "pf", "pmvg")}
    for tipo in ("pf", "pmvg"):
        d[tipo] = {c: f"{v:.2f}" for c, v in zip(ALIQUOTAS_ICMS, rec[tipo]) if v is not None}
    return json.dumps(d, sort_keys=True, ensure_ascii=False, separators=(",", ":"))


def _so_zero(valores: list) -> bool:
    """Preco publicado so a 0% (ou 0% e 'sem impostos') — tipico de isento de ICMS."""
    if valores[INDICE_ALIQUOTA["0"]] is None:
        return False
    return all(v is None for c, v in zip(ALIQUOTAS_ICMS, valores) if c not in ("SEM_IMPOSTOS", "0"))


def ler_lista(caminho: str) -> Lista:
    """Le e depura uma planilha. Nunca levanta: erro vai em Lista.erro."""
    t0 = time.time()
    lista = Lista(arquivo=caminho, nome=Path(caminho).name)
    try:
        _ler(lista)
    except LayoutError as e:
        lista.erro = str(e)
    except Exception as e:  # noqa: BLE001 — o relatorio final mostra o erro do arquivo
        lista.erro = f"{type(e).__name__}: {e}"
    lista.segundos = time.time() - t0
    return lista


def _ler(lista: Lista) -> None:
    from openpyxl import load_workbook

    lista.sequencia, lista.competencia, aba_esperada = competencia_do_nome(lista.nome)
    lista.sha256 = sha256_arquivo(lista.arquivo)
    st = lista.stats

    wb = load_workbook(lista.arquivo, read_only=True, data_only=True)
    try:
        if aba_esperada in wb.sheetnames:
            lista.aba = aba_esperada
        elif len(wb.sheetnames) == 1:
            lista.aba = wb.sheetnames[0]
            lista.avisos.append(f"aba '{lista.aba}' (esperado '{aba_esperada}') — unica aba, usada")
        else:
            raise LayoutError(f"aba '{aba_esperada}' nao encontrada entre {wb.sheetnames}")
        rows = wb[lista.aba].iter_rows(values_only=True)

        header = None
        for i, row in enumerate(rows, start=1):
            if i > LINHAS_CABECALHO:
                break
            ks = [chave(v) for v in row]
            if "REGISTRO" in ks and any(re.match(r"^(PF|PMVG)\b", k) for k in ks):
                header, lista.linha_cabecalho = row, i
                break
            if lista.declaradas is None:  # D1 = contagem de linhas de dados
                num = next((v for v in row if isinstance(v, (int, float)) and not isinstance(v, bool)), None)
                if num is not None:
                    lista.declaradas = int(num)
        if header is None:
            raise LayoutError(f"cabecalho (REGISTRO + PF/PMVG) nao encontrado nas {LINHAS_CABECALHO} primeiras linhas")

        precos: list[tuple[int, str, str]] = []
        campos: dict[str, int] = {}
        extras_cols: list[tuple[int, str]] = []
        vistos: set[tuple[str, str]] = set()
        com_cabecalho: set[int] = set()
        for j, h in enumerate(header):
            if chave(h) == "":
                continue
            com_cabecalho.add(j)
            cab = re.sub(r"\s+", " ", str(h)).strip()
            cp = coluna_preco(h)
            if cp:
                tipo, codigo, nota = cp
                if (tipo, codigo) in vistos:
                    raise LayoutError(f"coluna de preco repetida: {cab!r}")
                vistos.add((tipo, codigo))
                precos.append((j, tipo, codigo))
                destino = f"{tipo}:{codigo}" + (f" (nota de rodape {nota!r} no cabecalho)" if nota else "")
                lista.colunas.append({"col": _letra(j), "cabecalho": cab, "campo": destino})
                continue
            tipo_campo, nome = campo_da_coluna(h)
            if tipo_campo == "campo":
                if nome in campos:
                    raise LayoutError(f"duas colunas para o campo {nome!r}: {_letra(campos[nome])} e {_letra(j)}")
                campos[nome] = j
                lista.colunas.append({"col": _letra(j), "cabecalho": cab, "campo": nome})
            else:
                extras_cols.append((j, nome))
                lista.colunas.append({"col": _letra(j), "cabecalho": cab, "campo": f"extras.{nome}"})
        if "registro" not in campos:
            raise LayoutError("coluna REGISTRO nao encontrada")
        sem_cabecalho = [j for j in range(len(header)) if j not in com_cabecalho]

        def g(row: tuple, campo: str):
            j = campos.get(campo)
            return row[j] if j is not None and j < len(row) else None

        recs: list[dict] = []
        for row in rows:
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
                st["linhas_vazias"] += 1
                continue
            lista.linhas += 1
            for j in sem_cabecalho:
                if j < len(row) and texto(row[j]) is not None:
                    st[f"dado_sem_cabecalho_{_letra(j)}"] += 1

            reg_txt = texto(g(row, "registro")) or ""
            registro = re.sub(r"\D", "", reg_txt)
            if not registro:
                st["sem_registro"] += 1
                lista.exemplo("sem_registro", f"{reg_txt or '(vazio)'} | {texto(g(row, 'produto'))}")
                continue
            if len(registro) != 13:
                st["registro_nao13"] += 1
                lista.exemplo("registro_nao13", f"{registro} | {texto(g(row, 'produto'))}")

            pf: list = [None] * len(ALIQUOTAS_ICMS)
            pmvg: list = [None] * len(ALIQUOTAS_ICMS)
            notas: dict[str, dict[str, str]] = {"pf": {}, "pmvg": {}}
            liberado = False
            for j, tipo, codigo in precos:
                v = row[j] if j < len(row) else None
                valor, estado, nota = preco(v)
                if estado == "liberado":
                    liberado = True
                    st["celulas_liberado"] += 1
                elif estado == "invalido":
                    st["valor_invalido"] += 1
                    lista.exemplo("valor_invalido", f"{registro} {tipo.upper()} {codigo}: {v!r}")
                elif estado == "ambiguo":
                    st["valor_ambiguo"] += 1
                    lista.exemplo("valor_ambiguo", f"{registro} {tipo.upper()} {codigo}: {v!r}")
                if valor is not None:
                    (pf if tipo == "pf" else pmvg)[INDICE_ALIQUOTA[codigo]] = valor
                    if valor == 0:
                        st["preco_zero"] += 1
                    if nota:
                        notas[tipo][codigo] = nota

            extras: dict = {}
            for tipo in ("pf", "pmvg"):
                if not notas[tipo]:
                    continue
                arr = pf if tipo == "pf" else pmvg
                com_icms = {c for c, v in zip(ALIQUOTAS_ICMS, arr) if v is not None and c not in ("SEM_IMPOSTOS", "0")}
                marcas = set(notas[tipo].values())
                sufixo = "Pf" if tipo == "pf" else "Pmvg"
                if len(marcas) == 1 and set(notas[tipo]) == com_icms:
                    # caso normal: a mesma marca em todas as colunas com ICMS > 0
                    extras[f"nota{sufixo}"] = marcas.pop()
                else:
                    extras[f"notas{sufixo}"] = notas[tipo]
                    st["nota_irregular"] += 1
                    lista.exemplo("nota_irregular", f"{registro} {tipo.upper()}: {notas[tipo]}")
            if notas["pf"] or notas["pmvg"]:
                st["linhas_com_nota"] += 1

            rec = {"registro": registro, "precoLiberado": liberado, "pf": pf, "pmvg": pmvg}
            for c in TEXTOS:
                rec[c] = texto(g(row, c))
            eans = []
            for c in ("ean1", "ean2", "ean3"):
                e = texto(g(row, c))
                if e and e not in eans:
                    eans.append(e)
                    if not e.isdigit():
                        st["ean_irregular"] += 1
                        lista.exemplo("ean_irregular", f"{registro}: {e!r}")
            rec["ean"] = eans
            for c in BOOLEANOS:
                valor, reconhecido = booleano(g(row, c))
                rec[c] = valor
                if not reconhecido:
                    st["booleano_inesperado"] += 1
                    lista.exemplo("booleano_inesperado", f"{c}={g(row, c)!r}")
            # A partir de 2020 a coluna traz "(AR)" (analise recursal) ou notas de
            # rodape "(n)" (ex.: preco ajustado por decisao judicial; em fev/2024,
            # "(3)" em quase todas as linhas). So "(AR)"/Sim viram true; a marca
            # que nao e sim/nao fica em extras.analiseRecursal, e o booleano, null.
            marca = texto(g(row, "analiseRecursal"))
            valor, reconhecido = booleano(marca)
            if chave(marca) == "(AR)":
                rec["analiseRecursal"] = True
            elif reconhecido:
                rec["analiseRecursal"] = valor
            else:
                rec["analiseRecursal"] = None
                extras["analiseRecursal"] = marca
            for j, k in extras_cols:
                v = texto(row[j] if j < len(row) else None)
                if v is not None:
                    extras[k] = v
                    lista.extras[k] += 1
            rec["extras"] = extras
            recs.append(rec)
    finally:
        wb.close()

    if lista.declaradas is not None and lista.declaradas != lista.linhas:
        lista.avisos.append(f"D1 declara {lista.declaradas:,} linhas; lidas {lista.linhas:,}")

    # Duplicatas no mesmo mes, chave (registro, GGREM): identicas -> descarta;
    # so textos/booleanos diferentes -> mescla; precos diferentes -> CONFLITO
    # (mantem as duas versoes).
    grupos: dict[tuple, list[dict]] = {}
    for rec in recs:
        grupos.setdefault((rec["registro"], rec["codigoGgrem"]), []).append(rec)
    finais: list[dict] = []
    for (registro, ggrem), grupo in grupos.items():
        if len(grupo) == 1:
            finais.append(grupo[0])
            continue
        partes: dict[tuple, list[dict]] = {}
        for rec in grupo:
            partes.setdefault((tuple(rec["pf"]), tuple(rec["pmvg"]), rec["precoLiberado"]), []).append(rec)
        for parte in partes.values():
            base = parte[0]
            for outro in parte[1:]:
                if outro == base:
                    st["dup_identicas"] += 1
                else:
                    base = _mesclar(base, outro)
                    st["dup_mescladas"] += 1
            finais.append(base)
        if any(len(p) > 1 for p in partes.values()):
            lista.exemplo("duplicatas", f"{registro} (GGREM {ggrem}) {texto(grupo[0]['produto'])}: {len(grupo)} linhas -> {len(partes)}")
        if len(partes) > 1:
            st["dup_conflito"] += len(partes) - 1
            lista.exemplo("conflito", f"{registro} (GGREM {ggrem}) {texto(grupo[0]['produto'])}: {len(partes)} precos diferentes")

    presentes: set[tuple[str, str]] = set()
    caps: list[float] = []
    i0 = INDICE_ALIQUOTA["0"]
    for rec in finais:
        pf, pmvg = rec["pf"], rec["pmvg"]
        for c, a, b in zip(ALIQUOTAS_ICMS, pf, pmvg):
            if a is not None:
                presentes.add(("pf", c))
            if b is not None:
                presentes.add(("pmvg", c))
        st["liberados"] += rec["precoLiberado"]
        st["pf_so_0"] += _so_zero(pf)
        st["pmvg_so_0"] += _so_zero(pmvg)
        sem_pf = all(v is None for v in pf)
        st["cap_sem_pf"] += bool(rec["cap"]) and sem_pf
        st["sem_preco"] += sem_pf and all(v is None for v in pmvg) and not rec["precoLiberado"]
        if pf[i0] and pmvg[i0]:
            caps.append((1 - pmvg[i0] / pf[i0]) * 100)

        canon = _canonico(rec)
        h = hashlib.sha1(canon.encode("utf-8")).hexdigest()
        extras = json.dumps(rec["extras"], sort_keys=True, ensure_ascii=False) if rec["extras"] else None
        lista.versoes.append((
            h, rec["registro"], rec["codigoGgrem"], rec["ean"], rec["substancia"], rec["cnpj"],
            rec["laboratorio"], rec["produto"], rec["apresentacao"], rec["classeTerapeutica"],
            rec["tipoProduto"], rec["regimePreco"], rec["tarja"], rec["comercializacao"],
            rec["restricaoHospitalar"], rec["cap"], rec["confaz87"], rec["icms0"],
            rec["analiseRecursal"], rec["precoLiberado"],
            json.dumps(pf, separators=(",", ":")), json.dumps(pmvg, separators=(",", ":")), extras,
        ))
        lista.tamanhos.append(len(canon))
        lista.registros_set.add(rec["registro"])

    lista.registros = len(lista.registros_set)
    lista.aliquotas = [c for c in ALIQUOTAS_ICMS if ("pf", c) in presentes or ("pmvg", c) in presentes]
    lista.colunas_vazias = [
        f"{t.upper()} {rotulo_aliquota(codigo)}" for _, t, codigo in precos if (t, codigo) not in presentes
    ]
    lista.cap_percentual = round(statistics.median(caps), 2) if caps else None
    if len({v[0] for v in lista.versoes}) != len(lista.versoes):
        raise LayoutError("hash repetido depois da depuracao de duplicatas (bug do importador)")


def processar(arquivos: list[Path], jobs: int):
    """Le as planilhas em paralelo e devolve na ordem, com janela limitada
    (no maximo jobs + 1 listas lidas esperando a carga)."""
    if jobs <= 1:
        for a in arquivos:
            yield ler_lista(str(a))
        return
    with ProcessPoolExecutor(max_workers=jobs) as ex:
        fila: deque = deque()
        pendentes = iter(arquivos)
        for a in pendentes:
            fila.append(ex.submit(ler_lista, str(a)))
            if len(fila) > jobs:
                break
        while fila:
            lista = fila.popleft().result()
            proximo = next(pendentes, None)
            if proximo is not None:
                fila.append(ex.submit(ler_lista, str(proximo)))
            yield lista


# ─── conferencias opcionais ──────────────────────────────────────────────────


def ler_vigencia(caminho: str) -> dict[int, set[str]]:
    """Planilha de vigencia: linha 3 = n. da planilha, 4 = mes, 5 = ano (2 digitos);
    dados a partir da linha 6 (A = registro, B = total, C em diante = "V")."""
    from openpyxl import load_workbook

    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        rows = wb[wb.sheetnames[0]].iter_rows(values_only=True)
        topo = [next(rows, ()) for _ in range(5)]
        seqs: dict[int, int] = {}
        for j, v in enumerate(topo[2]):
            if j < 2 or v is None:
                continue
            try:
                seq = int(float(str(v).strip()))
            except ValueError:
                continue
            mes = chave(topo[3][j] if j < len(topo[3]) else "").lower()
            ano = texto(topo[4][j] if j < len(topo[4]) else None)
            if mes in MESES and ano and ano.isdigit():
                esperado = (2000 + int(ano) - 2017) * 12 + MESES.index(mes) + 1
                if esperado != seq:
                    print(f"  AVISO vigencia: coluna {_letra(j)} n. {seq:03d} diz {mes}/{ano} (esperado {esperado:03d})")
            seqs[j] = seq
        vig: dict[int, set[str]] = {s: set() for s in seqs.values()}
        for row in rows:
            registro = re.sub(r"\D", "", texto(row[0] if row else None) or "")
            if not registro:
                continue
            for j, seq in seqs.items():
                if j < len(row) and chave(row[j]) == "V":
                    vig[seq].add(registro)
    finally:
        wb.close()
    return {s: regs for s, regs in vig.items() if regs}


def ler_textos(caminho: str) -> dict[int, date]:
    """textos_CMED_original.docx: paragrafo 'NNN.mmm.aa' seguido da 1a data dd/mm/aaaa."""
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(caminho) as z:
        raiz = ElementTree.fromstring(z.read("word/document.xml"))
    datas: dict[int, date] = {}
    atual: int | None = None
    for p in raiz.iter(f"{w}p"):
        s = "".join(t.text or "" for t in p.iter(f"{w}t")).strip()
        m = re.fullmatch(r"(\d{3})\.([a-z]{3})\.(\d{2})", s, re.I)
        if m:
            atual = int(m.group(1))
            continue
        if atual is not None and atual not in datas:
            d = re.search(r"\b(\d{2})/(\d{2})/(\d{4})\b", s)
            if d:
                datas[atual] = date(int(d.group(3)), int(d.group(2)), int(d.group(1)))
    return datas


# ─── relatorio ───────────────────────────────────────────────────────────────


def _rotulos(codigos) -> str:
    return ", ".join(rotulo_aliquota(c) for c in codigos) or "(nenhuma)"


def relatorio(lista: Lista, novas: int | None, carga: dict | None) -> None:
    st = lista.stats
    comp = lista.competencia.strftime("%Y-%m") if lista.competencia else "?"
    print("")
    print(f"  {lista.nome}  competencia {comp}  aba '{lista.aba}'  cabecalho na linha {lista.linha_cabecalho}  ({lista.segundos:.1f}s)")
    if lista.erro:
        print(f"    ERRO: {lista.erro}")
        return
    decl = f"   (D1 declara {lista.declaradas:,})" if lista.declaradas is not None else ""
    print(f"    linhas lidas ................. {lista.linhas:,}{decl}")
    print(f"    descartadas sem registro ..... {st['sem_registro']:,}")
    print(f"    registros distintos .......... {lista.registros:,}")
    print(f"    versoes do mes ............... {len(lista.versoes):,}" + (f"   (novas: {novas:,})" if novas is not None else ""))
    print(f"    aliquotas presentes .......... {_rotulos(lista.aliquotas)}")
    if lista.colunas_vazias:
        print(f"    colunas de preco vazias ...... {', '.join(lista.colunas_vazias)}")
    extras = [c for c in lista.colunas if c["campo"].startswith("extras.")]
    if extras:
        desc = "; ".join(f"{c['cabecalho']} -> {c['campo']} ({lista.extras[c['campo'][7:]]:,})" for c in extras)
        print(f"    colunas nao mapeadas ......... {desc}")
    notas_cab = [c["cabecalho"] for c in lista.colunas if "nota de rodape" in c["campo"]]
    if notas_cab:
        print(f"    cabecalho com nota de rodape . {', '.join(notas_cab)}")
    print(f"    precos 'Liberado' ............ {st['liberados']:,} registros ({st['celulas_liberado']:,} celulas)")
    print(f"    so 0% publicado .............. PF {st['pf_so_0']:,} / PMVG {st['pmvg_so_0']:,}")
    print(f"    CAP = Sim sem PF ............. {st['cap_sem_pf']:,}")
    print(f"    sem nenhum preco ............. {st['sem_preco']:,}")
    print(f"    precos com nota (*) .......... {st['linhas_com_nota']:,} linhas ({st['nota_irregular']:,} com padrao irregular)")
    print(f"    valores invalidos ............ {st['valor_invalido']:,}   ambiguos: {st['valor_ambiguo']:,}   zerados: {st['preco_zero']:,}")
    print(f"    registros != 13 digitos ...... {st['registro_nao13']:,}")
    print(f"    duplicatas ................... identicas {st['dup_identicas']:,} / mescladas {st['dup_mescladas']:,} / conflito {st['dup_conflito']:,}")
    print(f"    CAP derivado ................. " + (f"{lista.cap_percentual:.2f}%".replace(".", ",") if lista.cap_percentual is not None else "-"))
    outros = {k: v for k, v in st.items() if k.startswith("dado_sem_cabecalho") or k in ("booleano_inesperado", "ean_irregular")}
    if outros:
        print(f"    outros ....................... {dict(outros)}")
    for aviso in lista.avisos:
        print(f"    AVISO: {aviso}")
    for tipo, valores in lista.exemplos.items():
        print(f"    ex. {tipo}: " + " || ".join(valores))
    if carga:
        print(
            f"    CARGA: {carga['inseridas']:,} versoes novas, {carga['reaproveitadas']:,} reaproveitadas; "
            f"vigencia {carga['vigencias']:,} (substituiu {carga['removidas']:,}) em {carga['segundos']:.1f}s"
        )
        print("    tabelas: " + "  ".join(f"{t} {carga['tamanhos'].get(t, 0) / 1e6:,.1f} MB" for t in TABELAS))


def conferir_vigencia(lista: Lista, vig: dict[int, set[str]]) -> int | None:
    marcas = vig.get(lista.sequencia)
    if marcas is None:
        print(f"    VIGENCIA: planilha de vigencia sem marcas para {lista.sequencia:03d}")
        return None
    so_lista = lista.registros_set - marcas
    so_vig = marcas - lista.registros_set
    dif = len(so_lista) + len(so_vig)
    print(f"    VIGENCIA: {len(marcas):,} marcas 'V'; diferenca simetrica {dif:,}" + ("  OK" if dif == 0 else ""))
    if so_lista:
        print(f"      so na lista ({len(so_lista):,}): {sorted(so_lista)[:MAX_EXEMPLOS]}")
    if so_vig:
        print(f"      so na vigencia ({len(so_vig):,}): {sorted(so_vig)[:MAX_EXEMPLOS]}")
    return dif


CSV_COLS = (
    "arquivo", "competencia", "status", "linhas", "declaradas", "sem_registro", "registros",
    "versoes", "novas", "aliquotas", "cap_percentual", "liberados", "pf_so_0", "pmvg_so_0",
    "cap_sem_pf", "sem_preco", "linhas_com_nota", "nota_irregular", "valor_invalido",
    "valor_ambiguo", "registro_nao13", "dup_identicas", "dup_mescladas", "dup_conflito",
    "dif_vigencia", "avisos",
)


def linha_csv(lista: Lista, novas, dif) -> dict:
    st = lista.stats
    return {
        "arquivo": lista.nome,
        "competencia": lista.competencia.isoformat() if lista.competencia else "",
        "status": f"ERRO: {lista.erro}" if lista.erro else "ok",
        "linhas": lista.linhas,
        "declaradas": lista.declaradas if lista.declaradas is not None else "",
        "registros": lista.registros,
        "versoes": len(lista.versoes),
        "novas": "" if novas is None else novas,
        "aliquotas": " ".join(lista.aliquotas),
        "cap_percentual": "" if lista.cap_percentual is None else f"{lista.cap_percentual:.2f}".replace(".", ","),
        "dif_vigencia": "" if dif is None else dif,
        "avisos": " | ".join(lista.avisos),
        **{k: st[k] for k in CSV_COLS if k in (
            "sem_registro", "liberados", "pf_so_0", "pmvg_so_0", "cap_sem_pf", "sem_preco",
            "linhas_com_nota", "nota_irregular", "valor_invalido", "valor_ambiguo",
            "registro_nao13", "dup_identicas", "dup_mescladas", "dup_conflito",
        )},
    }


# ─── banco ───────────────────────────────────────────────────────────────────


def conectar():
    import psycopg

    from import_redmine_base import load_database_url

    url = load_database_url()
    print(f"  banco ..... {re.sub(r'://[^@]*@', '://***@', url)}")
    return psycopg.connect(url, autocommit=True)


def tamanhos(cur) -> dict[str, int]:
    cur.execute(
        "SELECT c.relname, pg_total_relation_size(c.oid) FROM pg_class c"
        " JOIN pg_namespace n ON n.oid = c.relnamespace"
        " WHERE n.nspname = 'public' AND c.relkind = 'r' AND c.relname = ANY(%s)",
        (list(TABELAS),),
    )
    return dict(cur.fetchall())


def carregar(conn, lista: Lista, conhecidos: set[str], data_publicacao: date | None) -> dict:
    """Grava uma competencia numa transacao (REPLACE daquele mes)."""
    from psycopg.types.json import Jsonb

    t0 = time.time()
    hashes = [v[0] for v in lista.versoes]
    novas = [v for v in lista.versoes if v[0] not in conhecidos]
    cols = ", ".join(f'"{c}"' for c in VERSAO_COLS)
    tmp_ddl = ", ".join(f'"{c}" {TMP_TIPOS.get(c, "TEXT")}' for c in VERSAO_COLS)
    with conn.transaction(), conn.cursor() as cur:
        cur.execute("CREATE TEMP TABLE tmp_cmed_hash (hash TEXT PRIMARY KEY) ON COMMIT DROP")
        cur.execute(f"CREATE TEMP TABLE tmp_cmed_versao (ordem INTEGER, {tmp_ddl}) ON COMMIT DROP")
        with cur.copy("COPY tmp_cmed_hash (hash) FROM STDIN") as cp:
            for h in hashes:
                cp.write_row((h,))
        with cur.copy(f"COPY tmp_cmed_versao (ordem, {cols}) FROM STDIN") as cp:
            for i, v in enumerate(novas):
                cp.write_row((i, *v))
        cur.execute(
            f'INSERT INTO public."CmedPrecoVersao" ({cols}) SELECT {cols} FROM tmp_cmed_versao'
            ' ORDER BY ordem ON CONFLICT ("hash") DO NOTHING'
        )
        inseridas = cur.rowcount
        cur.execute('DELETE FROM public."CmedVigencia" WHERE "competencia" = %s', (lista.competencia,))
        removidas = cur.rowcount
        cur.execute(
            'INSERT INTO public."CmedVigencia" ("competencia", "versaoId")'
            ' SELECT %s, v."id" FROM tmp_cmed_hash t JOIN public."CmedPrecoVersao" v ON v."hash" = t.hash',
            (lista.competencia,),
        )
        vigencias = cur.rowcount
        if vigencias != len(hashes):
            raise RuntimeError(f"vigencia gravou {vigencias:,} de {len(hashes):,} versoes (hash ausente no banco)")
        cur.execute(
            """
            INSERT INTO public."CmedCompetencia"
                ("competencia", "sequencia", "dataPublicacao", "arquivo", "sha256", "linhas",
                 "registros", "aliquotas", "capPercentual", "colunasOriginais", "importadoEm")
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, CURRENT_TIMESTAMP)
            ON CONFLICT ("competencia") DO UPDATE SET
                "sequencia" = EXCLUDED."sequencia",
                "dataPublicacao" = COALESCE(EXCLUDED."dataPublicacao", "CmedCompetencia"."dataPublicacao"),
                "arquivo" = EXCLUDED."arquivo",
                "sha256" = EXCLUDED."sha256",
                "linhas" = EXCLUDED."linhas",
                "registros" = EXCLUDED."registros",
                "aliquotas" = EXCLUDED."aliquotas",
                "capPercentual" = EXCLUDED."capPercentual",
                "colunasOriginais" = EXCLUDED."colunasOriginais",
                "importadoEm" = CURRENT_TIMESTAMP
            """,
            (
                lista.competencia, lista.sequencia, data_publicacao, lista.nome, lista.sha256,
                lista.linhas, lista.registros, lista.aliquotas, lista.cap_percentual,
                Jsonb({"aba": lista.aba, "linhaCabecalho": lista.linha_cabecalho, "colunas": lista.colunas}),
            ),
        )
        tam = tamanhos(cur)
    conhecidos.update(hashes)
    return {
        "inseridas": inseridas,
        "reaproveitadas": len(hashes) - inseridas,
        "removidas": removidas,
        "vigencias": vigencias,
        "tamanhos": tam,
        "segundos": time.time() - t0,
    }


def limpar_orfas(cur, confirmar: bool) -> None:
    filtro = (
        ' FROM public."CmedPrecoVersao" v WHERE NOT EXISTS'
        ' (SELECT 1 FROM public."CmedVigencia" g WHERE g."versaoId" = v."id")'
    )
    cur.execute("SELECT count(*)" + filtro)
    n = cur.fetchone()[0]
    if not confirmar:
        print(f"  versoes sem vigencia: {n:,} (nada apagado no --dry-run)")
        return
    cur.execute('DELETE FROM public."CmedPrecoVersao" WHERE "id" IN (SELECT v."id"' + filtro + ")")
    print(f"  versoes sem vigencia apagadas: {cur.rowcount:,}")


# ─── principal ───────────────────────────────────────────────────────────────


def _intervalos(expr: str) -> set[int]:
    seqs: set[int] = set()
    for parte in expr.split(","):
        a, _, b = parte.strip().partition("-")
        if not a.isdigit() or (b and not b.isdigit()):
            raise SystemExit(f"ERRO: --somente invalido: {expr!r} (ex.: 037-048,050)")
        seqs.update(range(int(a), int(b or a) + 1))
    return seqs


def listar_arquivos(args) -> list[Path]:
    arquivos: list[Path] = []
    if args.dir:
        base = Path(args.dir)
        if not base.is_dir():
            raise SystemExit(f"ERRO: pasta nao encontrada: {base}")
        arquivos += [p for p in base.rglob("*.xlsx") if RE_ARQUIVO.match(p.name)]
    for f in args.file or []:
        p = Path(f)
        if not p.exists():
            raise SystemExit(f"ERRO: planilha nao encontrada: {p}")
        arquivos.append(p)
    por_seq: dict[int, Path] = {}
    for p in arquivos:
        m = RE_ARQUIVO.match(p.name)
        if not m:
            raise SystemExit(f"ERRO: nome fora do padrao NNN.mmm.aa.xlsx: {p.name}")
        seq = int(m.group(1))
        if seq in por_seq and por_seq[seq].resolve() != p.resolve():
            raise SystemExit(f"ERRO: duas planilhas para a sequencia {seq:03d}: {por_seq[seq]} e {p}")
        por_seq[seq] = p
    if args.somente:
        filtro = _intervalos(args.somente)
        por_seq = {s: p for s, p in por_seq.items() if s in filtro}
    return [por_seq[s] for s in sorted(por_seq)]


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    modo = ap.add_mutually_exclusive_group(required=True)
    modo.add_argument("--dry-run", action="store_true", help="le e valida sem tocar no banco")
    modo.add_argument("--confirm", action="store_true", help="grava no banco (REPLACE por competencia)")
    ap.add_argument("--dir", help="pasta com as planilhas NNN.mmm.aa.xlsx (busca em subpastas)")
    ap.add_argument("--file", action="append", help="planilha avulsa (pode repetir)")
    ap.add_argument("--somente", help="sequencias a processar, ex.: 037-048,050")
    ap.add_argument("--vigencia", help="planilha de vigencia para conferencia (diferenca tem de ser zero)")
    ap.add_argument("--textos", help="textos_CMED_original.docx: preenche dataPublicacao")
    ap.add_argument("--limpar-orfas", action="store_true", help="apaga versoes sem nenhuma vigencia")
    ap.add_argument("--recarregar", action="store_true", help="recarrega mesmo com sha256 igual ao gravado")
    ap.add_argument("--saida", help="grava o relatorio por arquivo em CSV (; e BOM, para o Excel)")
    ap.add_argument("--jobs", type=int, default=min(4, os.cpu_count() or 1), help="processos de leitura (padrao 4)")
    args = ap.parse_args()

    arquivos = listar_arquivos(args)
    if not arquivos and not args.textos and not args.limpar_orfas:
        ap.error("nenhuma planilha NNN.mmm.aa.xlsx: informe --dir ou --file")

    print("")
    print(f"  modo ...... {'DRY-RUN (nada e gravado)' if args.dry_run else 'CONFIRM (REPLACE por competencia)'}")
    if arquivos:
        print(f"  planilhas . {len(arquivos)} ({arquivos[0].name} .. {arquivos[-1].name}), leitura com {args.jobs} processo(s)")
    vig = None
    if args.vigencia:
        vig = ler_vigencia(args.vigencia)
        print(f"  vigencia .. {Path(args.vigencia).name}: {len(vig)} competencias com marcas")
    datas: dict[int, date] = {}
    if args.textos:
        datas = ler_textos(args.textos)
        print(f"  textos .... {Path(args.textos).name}: {len(datas)} datas de publicacao")

    conn = cur = None
    conhecidos: set[str] = set()
    gravadas: dict[int, str] = {}
    if args.confirm:
        conn = conectar()
        cur = conn.cursor()
        cur.execute(DDL)
        cur.execute("SELECT pg_database_size(current_database())")
        banco_antes = cur.fetchone()[0]
        cur.execute('SELECT "hash" FROM public."CmedPrecoVersao"')
        conhecidos = {r[0] for r in cur.fetchall()}
        cur.execute('SELECT "sequencia", "sha256" FROM public."CmedCompetencia"')
        gravadas = dict(cur.fetchall())
        print(f"  estado .... {len(gravadas)} competencias, {len(conhecidos):,} versoes; banco {banco_antes / 1e6:,.0f} MB")
        if datas and not arquivos:
            for seq, d in sorted(datas.items()):
                cur.execute('UPDATE public."CmedCompetencia" SET "dataPublicacao" = %s WHERE "sequencia" = %s', (d, seq))
            print(f"  dataPublicacao atualizada nas competencias carregadas ({len(datas)} datas lidas)")

    if not args.recarregar and gravadas:
        restantes = []
        for p in arquivos:
            seq = int(RE_ARQUIVO.match(p.name).group(1))
            if gravadas.get(seq) == sha256_arquivo(str(p)):
                print(f"  {p.name}: ja carregada com o mesmo sha256 — pulada (use --recarregar)")
            else:
                restantes.append(p)
        arquivos = restantes

    resumo: list[dict] = []
    vistos: set[str] = set()
    total_vig = total_novas_bytes = 0
    falhas = 0
    t0 = time.time()
    for lista in processar(arquivos, args.jobs):
        novas = dif = None
        carga = None
        if not lista.erro:
            hashes = [v[0] for v in lista.versoes]
            novas = sum(1 for h in hashes if h not in vistos and h not in conhecidos)
            total_novas_bytes += sum(t for v, t in zip(lista.versoes, lista.tamanhos) if v[0] not in vistos and v[0] not in conhecidos)
            vistos.update(hashes)
            total_vig += len(hashes)
            if conn is not None:
                try:
                    carga = carregar(conn, lista, conhecidos, datas.get(lista.sequencia))
                except Exception as e:  # noqa: BLE001 — falha num arquivo nao desfaz os anteriores
                    lista.erro = f"carga: {type(e).__name__}: {e}"
                    if conn.closed or conn.broken:
                        conn = conectar()
                        cur = conn.cursor()
        relatorio(lista, novas, carga)
        if vig is not None and not lista.erro:
            dif = conferir_vigencia(lista, vig)
        falhas += bool(lista.erro)
        resumo.append(linha_csv(lista, novas, dif))
        lista.versoes.clear()

    print("")
    print("  RESUMO")
    print(f"    planilhas ok ................. {len(resumo) - falhas} de {len(resumo)}")
    print(f"    versoes distintas ............ {len(vistos):,}" + (f" (alem das {len(conhecidos - vistos):,} ja gravadas)" if conhecidos else ""))
    print(f"    pares competencia x versao ... {total_vig:,}")
    print(f"    conteudo das versoes novas ... {total_novas_bytes / 1e6:,.1f} MB (JSON canonico)")
    print(f"    tempo ........................ {time.time() - t0:,.0f}s")
    for linha in resumo:
        if linha["status"] != "ok":
            print(f"    {linha['arquivo']}: {linha['status']}")

    if cur is not None:
        if args.limpar_orfas:
            limpar_orfas(cur, True)
        cur.execute('ANALYZE public."CmedCompetencia", public."CmedPrecoVersao", public."CmedVigencia"')
        tam = tamanhos(cur)
        cur.execute("SELECT pg_database_size(current_database())")
        print("    tabelas ...................... " + "  ".join(f"{t} {tam.get(t, 0) / 1e6:,.1f} MB" for t in TABELAS))
        print(f"    banco ........................ {cur.fetchone()[0] / 1e6:,.0f} MB")
        conn.close()
    elif args.limpar_orfas:
        conn = conectar()
        with conn:
            limpar_orfas(conn.cursor(), False)
    if args.dry_run:
        print("")
        print("  Nada foi gravado. Use --confirm para executar a carga.")

    if args.saida:
        with open(args.saida, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=CSV_COLS, delimiter=";")
            w.writeheader()
            w.writerows(resumo)
        print(f"  relatorio por arquivo: {args.saida}")
    return 1 if falhas else 0


if __name__ == "__main__":
    raise SystemExit(main())
